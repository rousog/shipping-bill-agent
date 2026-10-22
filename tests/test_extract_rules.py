from __future__ import annotations

import pytest

from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from shipping_agent.assets import LocalAssetStore
from shipping_agent.extractors.rules import extract_with_rules
from shipping_agent.loaders.eml import parse_eml
from shipping_agent.loaders.pdf_image import parse_pdf_image
from shipping_agent.loaders.router import route_path
from shipping_agent.models import RawDocument, SourceType
from shipping_agent.rules.registry import load_rules

SAMPLES = Path(__file__).resolve().parents[1] / "data" / "samples"

SAMPLES_AVAILABLE = SAMPLES.exists() and any(SAMPLES.iterdir())
pytestmark = pytest.mark.skipif(not SAMPLES_AVAILABLE, reason="confidential samples not present")


def test_extract_pdf2_line_items(tmp_path: Path) -> None:
    store = LocalAssetStore(tmp_path / "assets")
    doc = parse_pdf_image(SAMPLES / "PDF 示例账单2.pdf", store, use_ocr=False)
    result = extract_with_rules(doc, load_rules())

    assert result.document_type == "Debit Note"
    assert result.reference_no == "DC-108478"
    assert result.bl_no == "MAEU275752922"
    assert result.total_amount == Decimal("1410.40")
    assert len(result.charges) == 3
    assert {charge.amount for charge in result.charges} == {
        Decimal("65.00"),
        Decimal("1310.40"),
        Decimal("35.00"),
    }
    assert all(charge.currency == "USD" for charge in result.charges)
    # 规则层保留原始 P/C 值，标准化在 build_record 里做
    assert all(charge.payment_term == "C" for charge in result.charges)


def test_extract_pdf1_arrival_notice(tmp_path: Path) -> None:
    store = LocalAssetStore(tmp_path / "assets")
    doc = parse_pdf_image(SAMPLES / "PDF示例账单1.pdf", store, use_ocr=False)
    result = extract_with_rules(doc, load_rules())

    assert result.document_type == "Arrival Notice"
    assert result.reference_no == "CMDUCHN4053601"
    assert len(result.charges) == 1
    assert result.charges[0].amount == Decimal("245.0000")
    assert result.charges[0].currency == "USD"
    assert result.charges[0].payment_term == "C"


def test_extract_eml_body_charge(tmp_path: Path) -> None:
    store = LocalAssetStore(tmp_path / "assets")
    doc = parse_eml(next(SAMPLES.glob("*.eml")), store, use_ocr=False)
    result = extract_with_rules(doc, load_rules())

    assert result.document_type == "Pre-Alert"
    assert result.reference_no == "OOLU4055331281"
    assert len(result.charges) == 1
    assert result.charges[0].raw_name == "H/C"
    assert result.charges[0].amount == Decimal("50")
    assert result.charges[0].currency == "USD"
    assert result.charges[0].payment_term == "PP"


def test_extract_excel_table(tmp_path: Path) -> None:
    path = tmp_path / "charges.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["单据号", "费用名称", "金额", "币种", "付款方式"])
    sheet.append(["DC-XLSX-001", "HANDLING FEE", 65, "USD", "CC"])
    sheet.append(["DC-XLSX-001", "DOC FEE", 35, "USD", "CC"])
    workbook.save(path)

    store = LocalAssetStore(tmp_path / "assets")
    doc = route_path(path, store)
    result = extract_with_rules(doc, load_rules())

    assert result.reference_no == "DC-XLSX-001"
    assert len(result.charges) == 2
    assert all(charge.currency == "USD" for charge in result.charges)


def test_extract_word_table(tmp_path: Path) -> None:
    path = tmp_path / "bill.docx"
    document = Document()
    document.add_paragraph("DEBIT NOTE DC-DOCX-001")
    table = document.add_table(rows=2, cols=4)
    for cell, value in zip(
        table.rows[0].cells, ["费用名称", "金额", "币种", "付款方式"]
    ):
        cell.text = value
    for cell, value in zip(
        table.rows[1].cells, ["HANDLING FEE", "65.00", "USD", "CC"]
    ):
        cell.text = value
    document.save(path)

    store = LocalAssetStore(tmp_path / "assets")
    doc = route_path(path, store)
    result = extract_with_rules(doc, load_rules())

    assert result.document_type == "Debit Note"
    assert result.reference_no == "DC-DOCX-001"
    assert len(result.charges) == 1
    assert result.charges[0].currency == "USD"
    assert result.charges[0].payment_term == "CC"


def test_extract_image_uses_ocr_charges() -> None:
    doc = RawDocument(
        doc_id="IMG-TEST",
        asset_id="asset-test",
        source_type=SourceType.IMAGE,
        text=None,
        metadata={
            "ocr_used": True,
            "has_text_layer": False,
            "ocr": {
                "charges": [
                    {
                        "raw_name": "WHARFAGE",
                        "amount": "21.90",
                        "currency": None,
                        "payment_term": None,
                    }
                ]
            },
        },
    )
    result = extract_with_rules(doc, load_rules())
    assert len(result.charges) == 1
    assert result.charges[0].raw_name == "WHARFAGE"
    assert result.charges[0].amount == Decimal("21.90")
