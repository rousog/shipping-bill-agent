from __future__ import annotations

from pathlib import Path

from docx import Document

from shipping_agent.assets import LocalAssetStore
from shipping_agent.loaders.router import route_path


def test_parse_docx_with_table(tmp_path: Path) -> None:
    path = tmp_path / "bill.docx"
    document = Document()
    document.add_paragraph("DEBIT NOTE DC-001")
    document.add_paragraph("Issuer: Pilot Aiship")
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "费用名称"
    table.rows[0].cells[1].text = "金额"
    table.rows[0].cells[2].text = "币种"
    table.rows[1].cells[0].text = "HANDLING FEE"
    table.rows[1].cells[1].text = "65.00"
    table.rows[1].cells[2].text = "USD"
    document.save(path)

    store = LocalAssetStore(tmp_path / "assets")
    doc = route_path(path, store)

    assert "DEBIT NOTE" in (doc.text or "")
    assert "HANDLING FEE" in (doc.text or "")
    assert len(doc.tables) == 1
    assert doc.metadata["table_count"] == 1
