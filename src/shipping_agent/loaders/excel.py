"""Excel 解析：读取工作表、表头、费用明细行。

支持：
- .xlsx / .xlsm：openpyxl
- .xls：xlrd
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from ..assets import AssetStore
from ..models import AssetRef, RawDocument, SourceType

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "reference_no": (
        "单据号",
        "单号",
        "账单号",
        "发票号",
        "invoice no",
        "invoice",
        "d/c no",
        "dc no",
        "reference",
        "reference no",
        "document no",
    ),
    "bl_no": (
        "提单号",
        "提单",
        "b/l",
        "bl no",
        "bill of lading",
        "mbl",
        "hbl",
    ),
    "booking_no": ("订舱号", "订舱", "booking", "booking no"),
    "container_no": (
        "箱号",
        "集装箱号",
        "container",
        "container no",
        "cntr",
    ),
    "charge_item": (
        "费用名称",
        "费用",
        "收费项目",
        "项目",
        "charge",
        "charge item",
        "description",
        "item",
    ),
    "amount": ("金额", "合计", "总额", "amount", "total", "debit", "credit"),
    "currency": ("币种", "货币", "currency", "curr", "ccy"),
    "payment_term": (
        "付款方式",
        "预付到付",
        "p/c",
        "pp/cc",
        "payment",
        "prepaid/collect",
    ),
    "quantity": ("数量", "qty", "quantity"),
    "unit_price": ("单价", "unit price", "rate"),
}


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _normalize(value: str) -> str:
    value = value.lower().strip()
    value = re.sub(r"[\s_\-/]+", " ", value)
    return value


def _match_field(cell: str) -> str | None:
    normalized = _normalize(cell)
    if not normalized:
        return None

    # 先精确匹配，避免「提单号」被「单号」抢先匹配
    for field, aliases in _HEADER_ALIASES.items():
        if any(_normalize(alias) == normalized for alias in aliases):
            return field

    # 再按最长别名做包含匹配
    candidates: list[tuple[int, str]] = []
    for field, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            normalized_alias = _normalize(alias)
            if normalized_alias and normalized_alias in normalized:
                candidates.append((len(normalized_alias), field))

    if candidates:
        candidates.sort(reverse=True)
        return candidates[0][1]
    return None


def _score_header(row: list[str]) -> int:
    return len({field for cell in row if (field := _match_field(cell))})


def _map_header(row: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, cell in enumerate(row):
        field = _match_field(cell)
        if field and field not in mapping:
            mapping[field] = index
    return mapping


def _detect_header(rows: list[list[str]]) -> tuple[int, dict[str, int]]:
    best_index = 0
    best_mapping: dict[str, int] = {}
    best_score = 0
    for index, row in enumerate(rows[:15]):
        score = _score_header(row)
        if score > best_score:
            best_index = index
            best_mapping = _map_header(row)
            best_score = score
    return best_index, best_mapping


def _read_xlsx(path: Path) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    sheets: list[tuple[str, list[list[str]]]] = []
    workbook = load_workbook(path, data_only=True, read_only=True)
    for sheet in workbook.worksheets:
        rows = [
            [_cell_to_str(cell.value) for cell in row]
            for row in sheet.iter_rows()
        ]
        rows = [row for row in rows if any(row)]
        sheets.append((sheet.title, rows))
    workbook.close()
    return sheets


def _read_xls(path: Path) -> list[tuple[str, list[list[str]]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "读取 .xls 需要安装 xlrd：pip install xlrd"
        ) from exc

    sheets: list[tuple[str, list[list[str]]]] = []
    workbook = xlrd.open_workbook(str(path))
    for sheet in workbook.sheets():
        rows = [
            [
                _cell_to_str(sheet.cell_value(row_index, col_index))
                for col_index in range(sheet.ncols)
            ]
            for row_index in range(sheet.nrows)
        ]
        rows = [row for row in rows if any(row)]
        sheets.append((sheet.name, rows))
    return sheets


def _rows_to_text(sheet_name: str, rows: list[list[str]]) -> str:
    lines = [f"# Sheet: {sheet_name}"]
    for row in rows:
        lines.append(" | ".join(cell for cell in row if cell))
    return "\n".join(lines)


def parse_excel(
    source: str | Path | AssetRef,
    asset_store: AssetStore,
) -> RawDocument:
    asset = source if isinstance(source, AssetRef) else asset_store.put(source)
    path = asset_store.path(asset.asset_id)
    suffix = path.suffix.lower()

    if suffix == ".xls":
        sheets = _read_xls(path)
    elif suffix in {".xlsx", ".xlsm"}:
        sheets = _read_xlsx(path)
    else:
        raise ValueError(f"不支持的 Excel 后缀: {suffix}")

    tables: list[list[list[str]]] = []
    texts: list[str] = []
    metadata_sheets: list[dict] = []
    notes: list[str] = []
    total_rows = 0

    for sheet_name, rows in sheets:
        if not rows:
            metadata_sheets.append({"name": sheet_name, "rows": 0})
            continue
        header_index, mapping = _detect_header(rows)
        if not mapping:
            notes.append(f"工作表 {sheet_name} 未识别到表头")
        tables.append(rows)
        texts.append(_rows_to_text(sheet_name, rows))
        total_rows += len(rows)
        metadata_sheets.append(
            {
                "name": sheet_name,
                "rows": len(rows),
                "header_row": header_index,
                "column_mapping": mapping,
            }
        )

    metadata = {
        "file_name": asset.file_name,
        "mime_type": asset.mime_type,
        "size_bytes": asset.size_bytes,
        "sheet_count": len(sheets),
        "row_count": total_rows,
        "sheets": metadata_sheets,
    }

    return RawDocument(
        doc_id=f"XLSX-{asset.sha256[:12]}",
        asset_id=asset.asset_id,
        source_type=SourceType.EXCEL,
        text="\n\n".join(texts) or None,
        tables=tables,
        metadata=metadata,
        parse_notes=notes,
    )
