"""生成 Excel / Word 合成测试数据。

样例里没有 Excel、Word，Phase 1 用这些文件验证 loader。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

OUT_DIR = (
    Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "synthetic"
)


def make_excel() -> Path:
    path = OUT_DIR / "charges.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Charges"
    sheet.append(["单据号", "提单号", "费用名称", "金额", "币种", "付款方式"])
    sheet.append(["DC-SYN-001", "BL-SYN-001", "HANDLING FEE", 65, "USD", "CC"])
    sheet.append(["DC-SYN-001", "BL-SYN-001", "DOC FEE", 35, "USD", "CC"])
    workbook.save(path)
    return path


def make_word() -> Path:
    path = OUT_DIR / "debit_note.docx"
    document = Document()
    document.add_heading("DEBIT NOTE", level=1)
    document.add_paragraph("D/C No.: DC-SYN-002")
    document.add_paragraph("Issuer: Pilot Aiship Tech co., Ltd.")
    document.add_paragraph("B/L No.: BL-SYN-002")
    table = document.add_table(rows=2, cols=4)
    header = table.rows[0].cells
    header[0].text = "费用名称"
    header[1].text = "金额"
    header[2].text = "币种"
    header[3].text = "付款方式"
    row = table.rows[1].cells
    row[0].text = "HANDLING FEE"
    row[1].text = "65.00"
    row[2].text = "USD"
    row[3].text = "CC"
    document.save(path)
    return path


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    print("Excel:", make_excel())
    print("Word:", make_word())


if __name__ == "__main__":
    main()
